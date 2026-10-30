import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import IntegrityError
from django.db.models import Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse

from apps.directorio.forms import MascotaForm, VacunacionForm
from apps.usuarios.permissions import rol_requerido
from .models import Dueno, Mascota


@login_required
@rol_requerido('administrador', 'veterinario', 'recepcionista')
def registrar_dueno(request):
    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()
        telefono = request.POST.get('telefono', '').strip()
        direccion = request.POST.get('direccion', '').strip()
        correo = request.POST.get('correo', '').strip()

        try:
            dueno = Dueno(
                nombre_completo=nombre,
                telefono=telefono,
                direccion=direccion,
            )
            if hasattr(dueno, 'correo'):
                dueno.correo = correo
            dueno.save()

            messages.success(request, '¡Dueño registrado con éxito!')
            return redirect('registrar_dueno')

        except IntegrityError:
            messages.error(
                request,
                f"Error [ERR-001]: El número de teléfono '{telefono}' ya está registrado. "
                "Causa: el teléfono debe ser único. Solución: verifica el número o usa uno diferente."
            )
            return render(request, 'directorio/dueno_form.html', {
                'nombre': nombre,
                'telefono': telefono,
                'direccion': direccion,
                'correo': correo,
            })

    return render(request, 'directorio/dueno_form.html')


@login_required
@rol_requerido('administrador', 'veterinario', 'recepcionista')
def editar_dueno(request, dueno_id):
    """
    SW-02 / RF04 / PB-06: Edición de propietarios.
    Actualiza nombre, teléfono, dirección y correo verificando unicidad de teléfono.
    """
    dueno = get_object_or_404(Dueno, pk=dueno_id)

    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()
        telefono = request.POST.get('telefono', '').strip()
        direccion = request.POST.get('direccion', '').strip()
        correo = request.POST.get('correo', '').strip()

        try:
            dueno.nombre_completo = nombre
            dueno.telefono = telefono
            dueno.direccion = direccion
            if hasattr(dueno, 'correo'):
                dueno.correo = correo
            dueno.save()

            messages.success(request, '¡Propietario actualizado con éxito!')
            return redirect('buscar_directorio')

        except IntegrityError:
            messages.error(
                request,
                f"Error: El número de teléfono '{telefono}' ya está registrado con otro propietario. Debe ser único."
            )
            return render(request, 'directorio/dueno_form.html', {
                'dueno': dueno,
                'nombre': nombre,
                'telefono': telefono,
                'direccion': direccion,
                'correo': correo,
                'es_edicion': True,
            })

    return render(request, 'directorio/dueno_form.html', {
        'dueno': dueno,
        'nombre': dueno.nombre_completo,
        'telefono': dueno.telefono,
        'direccion': dueno.direccion,
        'correo': getattr(dueno, 'correo', ''),
        'es_edicion': True,
    })


@login_required
@rol_requerido('administrador', 'veterinario', 'recepcionista')
def lista_duenos(request):
    """Listado completo y numerado de dueños registrados."""
    duenos = Dueno.objects.all().order_by('id_dueno')
    return render(request, 'directorio/lista_duenos.html', {'duenos': duenos})


@login_required
@rol_requerido('administrador', 'veterinario', 'recepcionista')
def lista_mascotas(request):
    """Listado completo y numerado de mascotas registradas con su dueño."""
    mascotas = Mascota.objects.select_related('id_dueno').all().order_by('id_mascota')
    return render(request, 'directorio/lista_mascotas.html', {'mascotas': mascotas})


@login_required
@rol_requerido('administrador', 'veterinario', 'recepcionista')
def registrar_mascota(request):
    if request.method == 'POST':
        form = MascotaForm(request.POST)
        if form.is_valid():
            try:
                form.save()
            except IntegrityError:
                form.add_error(None, 'No se pudo guardar la mascota. El registro generado ya existe o hay un conflicto con los datos. Intenta nuevamente.')
            else:
                messages.success(request, '¡Mascota registrada con éxito!')
                return redirect('buscar_directorio')
    else:
        form = MascotaForm()

    return render(request, 'directorio/mascota_form.html', {'form': form})


@login_required
@rol_requerido('administrador', 'veterinario', 'recepcionista')
def editar_mascota(request, mascota_id):
    mascota = get_object_or_404(
        Mascota.objects.select_related('id_dueno'),
        id_mascota=mascota_id,
    )
    form = MascotaForm(instance=mascota)
    vacunacion_form = VacunacionForm()

    if request.method == 'POST' and 'guardar_vacunacion' in request.POST:
        vacunacion_form = VacunacionForm(request.POST)
        if vacunacion_form.is_valid():
            antecedente = vacunacion_form.save(commit=False)
            antecedente.mascota = mascota
            antecedente.save()
            messages.success(request, 'Antecedente de vacunación registrado.')
            return redirect('editar_mascota', mascota_id=mascota.id_mascota)
    elif request.method == 'POST':
        form = MascotaForm(request.POST, instance=mascota)
        if form.is_valid():
            try:
                form.save()
            except IntegrityError:
                form.add_error(
                    None,
                    'No se pudo actualizar la mascota por un conflicto con los datos. Intenta nuevamente.',
                )
            else:
                messages.success(request, '¡Mascota actualizada con éxito!')
                return redirect('editar_mascota', mascota_id=mascota.id_mascota)

    return render(request, 'directorio/mascota_form.html', {
        'form': form,
        'mascota': mascota,
        'vacunacion_form': vacunacion_form,
        'antecedentes_vacunacion': mascota.antecedentes_vacunacion.all(),
    })


@login_required
@rol_requerido('administrador', 'veterinario', 'recepcionista')
def buscar_directorio(request):
    """
    HU05 — "Como recepcionista, quiero buscar un dueño o mascota,
    para atenderlo rápidamente."

    Sirve tanto la página completa (para GET normal / sin JS) como
    respuestas JSON para la búsqueda en tiempo real vía fetch().
    """
    q = request.GET.get('q', '').strip()
    duenos = Dueno.objects.none()
    mascotas = Mascota.objects.none()

    if q:
        duenos = Dueno.objects.filter(
            Q(nombre_completo__icontains=q) | Q(telefono__icontains=q)
        ).order_by('nombre_completo')[:20]

        mascotas = Mascota.objects.select_related('id_dueno').filter(
            Q(nombre__icontains=q)
            | Q(id_dueno__nombre_completo__icontains=q)
            | Q(id_dueno__telefono__icontains=q)
        ).order_by('nombre')[:20]

    es_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest'
    if es_ajax:
        return JsonResponse({
            'duenos': [
                {
                    'id': d.id_dueno,
                    'nombre_completo': d.nombre_completo,
                    'telefono': d.telefono,
                    'direccion': d.direccion,
                    'url_edicion': reverse('editar_dueno', args=[d.id_dueno]),
                }
                for d in duenos
            ],
            'mascotas': [
                {
                    'id': m.id_mascota,
                    'nombre': m.nombre,
                    'especie': m.get_especie_display(),
                    'dueno_nombre': m.id_dueno.nombre_completo,
                    'dueno_telefono': m.id_dueno.telefono,
                    'url_edicion': reverse('editar_mascota', args=[m.id_mascota]),
                }
                for m in mascotas
            ],
        })

    return render(request, 'directorio/buscar.html', {'q': q, 'duenos': duenos, 'mascotas': mascotas})


@login_required
@rol_requerido('administrador', 'veterinario', 'recepcionista')
def autocomplete_duenos(request):
    q = request.GET.get('q', '').strip()
    resultados = []
    if q:
        duenos = Dueno.objects.filter(
            Q(nombre_completo__icontains=q) | Q(telefono__icontains=q)
        ).order_by('nombre_completo')[:15]
        for d in duenos:
            texto = f"{d.nombre_completo} ({d.telefono})" if d.telefono else d.nombre_completo
            resultados.append({
                'id': d.id_dueno,
                'texto': texto,
            })
    return JsonResponse({'resultados': resultados})


@login_required
@rol_requerido('administrador', 'veterinario', 'recepcionista')
def exportar_directorio_excel(request):
    """
    Genera y descarga un archivo Excel (.xlsx) con dos pestañas: Dueños y Mascotas.
    """
    wb = openpyxl.Workbook()

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E1B2E', end_color='1E1B2E', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')

    # Pestaña 1: Dueños
    ws_duenos = wb.active
    ws_duenos.title = "Dueños"

    headers_duenos = ["#", "ID Dueño", "Nombre Completo", "Teléfono", "Dirección", "Correo Electrónico"]
    ws_duenos.append(headers_duenos)

    for col_num, header in enumerate(headers_duenos, 1):
        cell = ws_duenos.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    duenos = Dueno.objects.all().order_by('id_dueno')
    for idx, d in enumerate(duenos, 1):
        correo = getattr(d, 'correo', '') or '-'
        ws_duenos.append([
            idx,
            d.id_dueno,
            d.nombre_completo,
            d.telefono,
            d.direccion or '-',
            correo,
        ])

    for col in ws_duenos.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_duenos.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Pestaña 2: Mascotas
    ws_mascotas = wb.create_sheet(title="Mascotas")

    headers_mascotas = ["#", "ID Mascota", "Nombre Mascota", "Especie", "Raza", "Propietario", "Teléfono Propietario"]
    ws_mascotas.append(headers_mascotas)

    for col_num, header in enumerate(headers_mascotas, 1):
        cell = ws_mascotas.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    mascotas = Mascota.objects.select_related('id_dueno').all().order_by('id_mascota')
    for idx, m in enumerate(mascotas, 1):
        especie_str = m.get_especie_display() if hasattr(m, 'get_especie_display') else str(m.especie)
        ws_mascotas.append([
            idx,
            m.id_mascota,
            m.nombre,
            especie_str,
            m.raza or '-',
            m.id_dueno.nombre_completo if m.id_dueno else '-',
            m.id_dueno.telefono if m.id_dueno else '-',
        ])

    for col in ws_mascotas.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_mascotas.column_dimensions[col_letter].width = max(max_len + 4, 12)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="Directorio_PradosPets.xlsx"'
    wb.save(response)
    return response


@login_required
@rol_requerido('administrador', 'veterinario', 'recepcionista')
def exportar_duenos_excel(request):
    """
    Genera y descarga un archivo Excel exclusivo para la lista de Dueños.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dueños"

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='6366F1', end_color='6366F1', fill_type='solid')

    headers = ["#", "ID Dueño", "Nombre Completo", "Teléfono", "Dirección", "Correo Electrónico"]
    ws.append(headers)

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill

    duenos = Dueno.objects.all().order_by('id_dueno')
    for idx, d in enumerate(duenos, 1):
        ws.append([
            idx,
            d.id_dueno,
            d.nombre_completo,
            d.telefono,
            d.direccion or '-',
            getattr(d, 'correo', '') or '-',
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 4, 12)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Listado_Duenos_PradosPets.xlsx"'
    wb.save(response)
    return response


@login_required
@rol_requerido('administrador', 'veterinario', 'recepcionista')
def exportar_mascotas_excel(request):
    """
    Genera y descarga un archivo Excel exclusivo para la lista de Mascotas.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mascotas"

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='10B981', end_color='10B981', fill_type='solid')

    headers = ["#", "ID Mascota", "Nombre Mascota", "Especie", "Raza", "Propietario", "Teléfono Propietario"]
    ws.append(headers)

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill

    mascotas = Mascota.objects.select_related('id_dueno').all().order_by('id_mascota')
    for idx, m in enumerate(mascotas, 1):
        ws.append([
            idx,
            m.id_mascota,
            m.nombre,
            m.get_especie_display() if hasattr(m, 'get_especie_display') else str(m.especie),
            m.raza or '-',
            m.id_dueno.nombre_completo if m.id_dueno else '-',
            m.id_dueno.telefono if m.id_dueno else '-',
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 4, 12)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Listado_Mascotas_PradosPets.xlsx"'
    wb.save(response)
    return response