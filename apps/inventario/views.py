from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.db import transaction, models
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse_lazy
from django.views.generic import ListView, CreateView

from apps.usuarios.permissions import rol_requerido
from .forms import ProductoForm, ReposicionForm, CargaExcelForm
from .models import Producto, Medicamento, Servicio, StockLote, MovimientoInventario, CategoriaProducto


class InventarioListView(LoginRequiredMixin, ListView):
    model = Producto
    template_name = 'inventario/inventario_list.html'
    context_object_name = 'productos'
    paginate_by = 15

    def get_queryset(self):
        queryset = Producto.objects.filter(activo=True).select_related('categoria', 'proveedor')
        q = self.request.GET.get('q')
        categoria_id = self.request.GET.get('categoria')
        
        if q:
            queryset = queryset.filter(
                models.Q(nombre__icontains=q) | 
                models.Q(codigo__icontains=q) | 
                models.Q(codigo_barras__icontains=q)
            )
        if categoria_id:
            queryset = queryset.filter(categoria_id=categoria_id)
            
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['categorias'] = CategoriaProducto.objects.all()
        productos = Producto.objects.filter(activo=True)
        context['total_valor_inventario'] = sum(p.valor_total_inventario for p in productos)
        context['productos_bajo_stock'] = productos.filter(stock_actual__lte=models.F('stock_minimo')).count()
        return context


class ProductoCreateView(LoginRequiredMixin, CreateView):
    model = Producto
    form_class = ProductoForm
    template_name = 'inventario/producto_form.html'
    success_url = reverse_lazy('inventario:lista')

    def form_valid(self, form):
        with transaction.atomic():
            producto = form.save()
            if producto.stock_actual > 0:
                MovimientoInventario.objects.create(
                    producto=producto,
                    tipo='ENTRADA_COMPRA',
                    cantidad=producto.stock_actual,
                    precio_unitario_historico=producto.precio_compra,
                    monto_total=producto.stock_actual * producto.precio_compra,
                    usuario=self.request.user,
                    observacion='Inventario Inicial / Registro de Producto'
                )
        messages.success(self.request, "Producto creado exitosamente con su movimiento contable inicial.")
        return redirect(self.success_url)


@login_required
def registrar_reposicion(request):
    if request.method == 'POST':
        form = ReposicionForm(request.POST)
        if form.is_valid():
            producto = form.cleaned_data['producto']
            cantidad = form.cleaned_data['cantidad']
            nuevo_costo = form.cleaned_data.get('nuevo_precio_compra')
            observacion = form.cleaned_data.get('observacion')

            with transaction.atomic():
                if nuevo_costo is not None:
                    producto.precio_compra = nuevo_costo
                
                producto.stock_actual += cantidad
                producto.save()

                costo_aplicado = nuevo_costo if nuevo_costo is not None else producto.precio_compra

                MovimientoInventario.objects.create(
                    producto=producto,
                    tipo='ENTRADA_COMPRA',
                    cantidad=cantidad,
                    precio_unitario_historico=costo_aplicado,
                    monto_total=cantidad * costo_aplicado,
                    usuario=request.user,
                    observacion=observacion or "Reabastecimiento de stock"
                )

            messages.success(request, f"Stock actualizado para {producto.nombre}. Nuevo stock: {producto.stock_actual}")
            return redirect('inventario:lista')
    else:
        form = ReposicionForm()
    
    return render(request, 'inventario/reposicion_form.html', {'form': form})


@login_required
@rol_requerido('administrador', 'veterinario', 'recepcionista')
def cargar_catalogo_excel(request):
    errores = []
    exitos = 0

    if request.method == 'POST':
        form = CargaExcelForm(request.POST, request.FILES)
        if form.is_valid():
            archivo = request.FILES['archivo_excel']
            
            try:
                wb = openpyxl.load_workbook(archivo, data_only=True)
                sheet = wb.active
            except Exception as e:
                messages.error(request, f"Error al abrir el archivo Excel: {str(e)}")
                return render(request, 'inventario/cargar_excel.html', {'form': form, 'errores': errores, 'exitos': exitos})

            for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if not any(row):
                    continue

                tipo = str(row[0] or '').strip().upper()
                codigo = str(row[1] or '').strip()
                nombre = str(row[2] or '').strip()
                descripcion = str(row[3] or '').strip()
                precio_raw = row[4]
                stock_min_raw = row[5]

                if not tipo or tipo not in ('PRODUCTO', 'MEDICAMENTO', 'SERVICIO'):
                    errores.append({'fila': row_idx, 'error': f"Tipo inválido ('{tipo}'). Debe ser PRODUCTO, MEDICAMENTO o SERVICIO."})
                    continue

                if not codigo:
                    errores.append({'fila': row_idx, 'error': "El campo 'Código' es obligatorio."})
                    continue

                if not nombre:
                    errores.append({'fila': row_idx, 'error': "El campo 'Nombre' es obligatorio."})
                    continue

                try:
                    precio = float(precio_raw)
                    if precio < 0:
                        raise ValueError
                except (ValueError, TypeError):
                    errores.append({'fila': row_idx, 'error': f"Precio inválido ('{precio_raw}'). Debe ser un número mayor o igual a 0."})
                    continue

                try:
                    with transaction.atomic():
                        if tipo == 'SERVICIO':
                            if Servicio.objects.filter(codigo=codigo).exists():
                                errores.append({'fila': row_idx, 'error': f"El servicio con código '{codigo}' ya existe."})
                                continue

                            Servicio.objects.create(
                                codigo=codigo,
                                nombre=nombre,
                                descripcion=descripcion,
                                precio=precio,
                            )
                            exitos += 1

                        elif tipo in ('PRODUCTO', 'MEDICAMENTO'):
                            if Producto.objects.filter(codigo=codigo).exists():
                                errores.append({'fila': row_idx, 'error': f"El producto/medicamento con código '{codigo}' ya existe."})
                                continue

                            try:
                                stock_min = int(stock_min_raw) if stock_min_raw is not None else 5
                            except (ValueError, TypeError):
                                stock_min = 5

                            producto = Producto.objects.create(
                                codigo=codigo,
                                nombre=nombre,
                                descripcion=descripcion,
                                precio_venta=precio,
                                stock_minimo=stock_min,
                            )

                            if tipo == 'MEDICAMENTO':
                                principio = str(row[6] or '').strip() if len(row) > 6 else ''
                                concentracion = str(row[7] or '').strip() if len(row) > 7 else ''
                                presentacion = str(row[8] or '').strip() if len(row) > 8 else ''
                                receta_val = str(row[9] or '').strip().upper() in ('SI', 'S', 'TRUE', '1') if len(row) > 9 else False

                                Medicamento.objects.create(
                                    producto=producto,
                                    principio_activo=principio,
                                    concentracion=concentracion,
                                    presentacion=presentacion,
                                    requiere_receta=receta_val,
                                )

                            num_lote = str(row[10] or '').strip() if len(row) > 10 else ''
                            fecha_venc_raw = row[11] if len(row) > 11 else None
                            cant_lote_raw = row[12] if len(row) > 12 else None

                            if num_lote and cant_lote_raw is not None:
                                cant_lote = int(cant_lote_raw)
                                fecha_venc = None
                                
                                if isinstance(fecha_venc_raw, datetime):
                                    fecha_venc = fecha_venc_raw.date()
                                elif isinstance(fecha_venc_raw, str) and fecha_venc_raw.strip():
                                    fecha_venc = datetime.strptime(fecha_venc_raw.strip(), "%Y-%m-%d").date()

                                StockLote.objects.create(
                                    producto=producto,
                                    numero_lote=num_lote,
                                    fecha_vencimiento=fecha_venc,
                                    cantidad=cant_lote,
                                )

                                producto.stock_actual += cant_lote
                                producto.save()

                                MovimientoInventario.objects.create(
                                    producto=producto,
                                    tipo='ENTRADA_COMPRA',
                                    cantidad=cant_lote,
                                    precio_unitario_historico=precio,
                                    monto_total=precio * cant_lote,
                                    usuario=request.user,
                                    observacion=f"Carga inicial Excel - Lote {num_lote}"
                                )

                            exitos += 1

                except Exception as ex:
                    if not any(e['fila'] == row_idx for e in errores):
                        errores.append({'fila': row_idx, 'error': f"Error al guardar registro: {str(ex)}"})

            if exitos > 0:
                messages.success(request, f"Se cargaron exitosamente {exitos} registros en el catálogo.")
            if not errores and exitos == 0:
                messages.warning(request, "El archivo no contenía datos válidos para procesar.")

    else:
        form = CargaExcelForm()

    return render(
        request,
        'inventario/cargar_excel.html',
        {'form': form, 'errores': errores, 'exitos': exitos}
    )


@login_required
def exportar_inventario_excel(request):
    """
    Genera un archivo .xlsx con el listado actual de productos en inventario.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventario Prados Pets"

    headers = [
        'Código', 'Código Barras', 'Nombre', 'Categoría', 'Proveedor',
        'Precio Compra ($)', 'Precio Venta ($)', 'Stock Actual', 'Stock Mínimo', 'Estado'
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1E1530", end_color="1E1530", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    productos = Producto.objects.filter(activo=True).select_related('categoria', 'proveedor')
    for prod in productos:
        estado = "REABASTECER" if prod.necesita_reposicion else "OK"
        cat_nombre = prod.categoria.nombre if prod.categoria else "General"
        prov_nombre = prod.proveedor.nombre if prod.proveedor else "N/A"
        
        ws.append([
            prod.codigo,
            prod.codigo_barras or '',
            prod.nombre,
            cat_nombre,
            prov_nombre,
            float(prod.precio_compra),
            float(prod.precio_venta),
            prod.stock_actual,
            prod.stock_minimo,
            estado
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="Inventario_Prados_Pets.xlsx"'
    wb.save(response)
    return response


@login_required
def descargar_plantilla_excel(request):
    """
    Genera un archivo .xlsx formateado con ejemplos para la carga masiva.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plantilla Catálogo"

    headers = [
        'Tipo', 'Codigo', 'Nombre', 'Descripcion', 'Precio', 'StockMinimo',
        'PrincipioActivo', 'Concentracion', 'Presentacion', 'RequiereReceta',
        'NumeroLote', 'FechaVencimiento', 'CantidadLote'
    ]
    ws.append(headers)

    ejemplos = [
        ['MEDICAMENTO', 'MED-001', 'Amoxicilina 500mg', 'Antibiótico de amplio espectro', 25000, 5, 'Amoxicilina', '500 mg', 'Frasco 100ml', 'SI', 'LOTE-2026-A', '2027-12-31', 30],
        ['PRODUCTO', 'PROD-001', 'Champú Antipulgas 250ml', 'Higiene y cuidado canino', 32000, 3, '', '', 'Frasco 250ml', 'NO', 'LOTE-2026-B', '2028-06-30', 15],
        ['SERVICIO', 'SERV-001', 'Consulta Médica General', 'Atención médica veterinaria integral', 45000, 0, '', '', '', 'NO', '', '', '']
    ]
    for row in ejemplos:
        ws.append(row)

    header_fill = PatternFill(start_color="1E1530", end_color="1E1530", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="Plantilla_Carga_Catalogo_PradosPets.xlsx"'
    wb.save(response)
    return response


class KardexListView(LoginRequiredMixin, ListView):
    model = MovimientoInventario
    template_name = 'inventario/kardex.html'
    context_object_name = 'movimientos'
    paginate_by = 20